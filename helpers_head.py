# -*- coding: utf-8 -*-
"""Builds the revised manuscript (full + blinded) as .docx following the promotion-track structure guide:
A4, Times New Roman 12 pt, 2.54 cm margins, numbered headings, APA 7 references, 'Page X of Y' footer."""
import json, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/private/tmp/claude-501/-Users-mouatafaelbehairy-Desktop-business/f52f105c-d03b-4099-ad6c-6372f7dda005/scratchpad"
R = json.load(open(f"{OUT}/results.json"))
BLIND = "--blind" in sys.argv
FONT = "Times New Roman"

def f4(x): return f"{x:.4f}"
def f3(x): return f"{x:.3f}"
def pct(x, d=1): return f"{100*x:.{d}f}%"

# ------------------------------------------------------------------ document setup
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(sec, side, Cm(2.54))
st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT); st.element.rPr.rFonts.set(qn("w:cs"), FONT)
pf = st.paragraph_format; pf.space_after = Pt(6); pf.line_spacing = 1.0
for lvl, size in ((1, 14), (2, 12), (3, 12)):
    h = doc.styles[f"Heading {lvl}"]; h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True
    h.font.italic = (lvl == 3); h.font.color.rgb = RGBColor(0, 0, 0)
    h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT); h.element.rPr.rFonts.set(qn("w:ascii"), FONT); h.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    h.paragraph_format.space_before = Pt(12 if lvl == 1 else 8); h.paragraph_format.space_after = Pt(4); h.paragraph_format.keep_with_next = True

def add_field(run, instr):
    for tag, txt in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if tag:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); run._r.append(e)
        else:
            e = OxmlElement("w:instrText") if txt == instr else OxmlElement("w:t")
            e.set(qn("xml:space"), "preserve"); e.text = txt; run._r.append(e)

fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Page "); r.font.size = Pt(9); r.font.name = FONT
r = fp.add_run(); r.font.size = Pt(9); add_field(r, "PAGE")
r = fp.add_run(" of "); r.font.size = Pt(9); r.font.name = FONT
r = fp.add_run(); r.font.size = Pt(9); add_field(r, "NUMPAGES")

def P(text="", bold=False, italic=False, align=None, size=None, after=None, before=None, indent=False, keep=False):
    p = doc.add_paragraph()
    if align == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "j" or align is None: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "r": p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "l": p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if before is not None: p.paragraph_format.space_before = Pt(before)
    if indent: p.paragraph_format.first_line_indent = Cm(0.75)
    if keep: p.paragraph_format.keep_with_next = True
    # simple inline markup: **bold**, *italic*
    import re
    for tok in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if not tok: continue
        if tok.startswith("**"): rr = p.add_run(tok[2:-2]); rr.bold = True
        elif tok.startswith("*"): rr = p.add_run(tok[1:-1]); rr.italic = True
        else: rr = p.add_run(tok)
        rr.font.name = FONT
        if bold: rr.bold = True
        if italic: rr.italic = True
        if size: rr.font.size = Pt(size)
    return p

def H(text, lvl=1): return doc.add_heading(text, lvl)

def RTL(p):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:bidi"); b.set(qn("w:val"), "1"); pPr.append(b)
    for rr in p.runs:
        rPr = rr._r.get_or_add_rPr(); rtl = OxmlElement("w:rtl"); rtl.set(qn("w:val"), "1"); rPr.append(rtl)
        rr.font.name = "Traditional Arabic"; rPr.rFonts.set(qn("w:cs"), "Traditional Arabic"); rr.font.size = Pt(13)

def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor); tcPr.append(shd)

def TABLE(caption, header, rows, note=None, widths=None, size=9.5):
    cp = P(caption, bold=False, align="l", after=3, keep=True)
    for rr in cp.runs: rr.font.size = Pt(11)
    # make "Table N." bold
    cp.runs[0].text = ""; import re
    m = re.match(r"(Table \d+\.)(.*)", caption, re.S)
    cp.clear(); r1 = cp.add_run(m.group(1)); r1.bold = True; r1.font.size = Pt(11); r1.font.name = FONT
    r2 = cp.add_run(m.group(2)); r2.italic = True; r2.font.size = Pt(11); r2.font.name = FONT
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; rr = c.paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(size); rr.font.name = FONT
        set_cell_bg(c, "E7EEF5")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; para = cells[i].paragraphs[0]; para.paragraph_format.space_after = Pt(0)
            bold = False
            if isinstance(v, tuple): v, bold = v
            rr = para.add_run(str(v)); rr.font.size = Pt(size); rr.font.name = FONT; rr.bold = bold
    if widths:
        t.autofit = False
        tblPr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        for i, w in enumerate(widths): t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    for row in t.rows:
        for c in row.cells:
            for para in c.paragraphs: para.paragraph_format.line_spacing = 1.0; para.paragraph_format.space_after = Pt(0)
    if note:
        np_ = P(note, align="l", size=9.5, after=10, before=2)
    else:
        P("", after=6)

def FIG(path, caption, width_cm=13.5):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width_cm))
    import re
    m = re.match(r"(Figure \d+\.)(.*)", caption, re.S)
    cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; cp.paragraph_format.space_after = Pt(10)
    r1 = cp.add_run(m.group(1)); r1.bold = True; r1.font.size = Pt(10.5); r1.font.name = FONT
    r2 = cp.add_run(m.group(2)); r2.font.size = Pt(10.5); r2.font.name = FONT

def EQ(text, num):
    t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    a, b = t.rows[0].cells; a.width = Cm(13.5); b.width = Cm(2.0)
    a.text = ""; b.text = ""
    ra = a.paragraphs[0].add_run(text); ra.font.name = "Cambria Math"; ra.font.size = Pt(12); a.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    rb = b.paragraphs[0].add_run(f"({num})"); rb.font.name = FONT; rb.font.size = Pt(12); b.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for c in (a, b): c.paragraphs[0].paragraph_format.space_after = Pt(4)

